import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

# =========================================================
# BOOK INVENTORY MANAGEMENT SYSTEM
# Pydroid 3 Optimized Version
# =========================================================

FILE_NAME = "book_inventory.xlsx"


# =========================================================
# DATABASE / EXCEL FUNCTIONS
# =========================================================

def create_database():
    """Create Excel file if it does not exist."""

    if not os.path.exists(FILE_NAME):
        wb = Workbook()
        ws = wb.active
        ws.title = "Books"

        ws.append([
            "Book ID",
            "Title",
            "Author",
            "Category",
            "Price",
            "Quantity"
        ])

        wb.save(FILE_NAME)


def load_books():
    """Load all books from Excel."""

    create_database()

    wb = load_workbook(FILE_NAME)
    ws = wb["Books"]

    books = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            books.append(row)

    wb.close()

    return books


def save_books(books):
    """Save book list to Excel."""

    wb = Workbook()
    ws = wb.active
    ws.title = "Books"

    ws.append([
        "Book ID",
        "Title",
        "Author",
        "Category",
        "Price",
        "Quantity"
    ])

    for book in books:
        ws.append(list(book))

    wb.save(FILE_NAME)


# =========================================================
# MAIN APPLICATION
# =========================================================

class BookInventoryApp:

    def __init__(self, root):

        self.root = root

        self.root.title("Book Inventory Management")

        # Phone-friendly size
        self.root.geometry("390x700")

        self.root.resizable(False, False)

        self.books = []

        create_database()

        self.show_login()


    # =====================================================
    # CLEAR SCREEN
    # =====================================================

    def clear_screen(self):

        for widget in self.root.winfo_children():
            widget.destroy()


    # =====================================================
    # LOGIN PAGE
    # =====================================================

    def show_login(self):

        self.clear_screen()

        frame = tk.Frame(self.root)
        frame.pack(expand=True)

        title = tk.Label(
            frame,
            text="BOOK INVENTORY",
            font=("Arial", 22, "bold")
        )
        title.pack(pady=10)

        subtitle = tk.Label(
            frame,
            text="Management System",
            font=("Arial", 14)
        )
        subtitle.pack(pady=5)

        tk.Label(
            frame,
            text="Username",
            font=("Arial", 12)
        ).pack(pady=(30, 5))

        self.username_entry = tk.Entry(
            frame,
            font=("Arial", 13),
            width=25
        )
        self.username_entry.pack()

        tk.Label(
            frame,
            text="Password",
            font=("Arial", 12)
        ).pack(pady=(15, 5))

        self.password_entry = tk.Entry(
            frame,
            font=("Arial", 13),
            width=25,
            show="*"
        )
        self.password_entry.pack()

        tk.Button(
            frame,
            text="LOGIN",
            font=("Arial", 13, "bold"),
            width=20,
            height=2,
            command=self.login
        ).pack(pady=30)

        tk.Label(
            frame,
            text="Demo Login\nUsername: admin\nPassword: admin123",
            font=("Arial", 10)
        ).pack()


    def login(self):

        username = self.username_entry.get()
        password = self.password_entry.get()

        if username == "admin" and password == "admin123":

            self.show_dashboard()

        else:

            messagebox.showerror(
                "Login Failed",
                "Invalid username or password."
            )


    # =====================================================
    # DASHBOARD
    # =====================================================

    def show_dashboard(self):

        self.clear_screen()

        self.books = load_books()

        title = tk.Label(
            self.root,
            text="BOOK INVENTORY",
            font=("Arial", 22, "bold")
        )
        title.pack(pady=20)

        total_books = len(self.books)

        total_quantity = 0

        for book in self.books:

            try:
                total_quantity += int(book[5])
            except:
                pass

        tk.Label(
            self.root,
            text="Total Different Books: " + str(total_books),
            font=("Arial", 14)
        ).pack(pady=5)

        tk.Label(
            self.root,
            text="Total Quantity: " + str(total_quantity),
            font=("Arial", 14)
        ).pack(pady=5)

        # Buttons

        buttons = [

            ("ADD BOOK", self.add_book),

            ("VIEW BOOKS", self.view_books),

            ("SEARCH BOOK", self.search_book),

            ("UPDATE BOOK", self.update_book),

            ("DELETE BOOK", self.delete_book),

            ("EXIT", self.root.destroy)

        ]

        for text, command in buttons:

            tk.Button(
                self.root,
                text=text,
                font=("Arial", 12, "bold"),
                width=25,
                height=2,
                command=command
            ).pack(pady=6)

        tk.Button(
            self.root,
            text="LOGOUT",
            font=("Arial", 11),
            width=15,
            command=self.show_login
        ).pack(pady=15)


    # =====================================================
    # ADD BOOK
    # =====================================================

    def add_book(self):

        self.clear_screen()

        tk.Label(
            self.root,
            text="ADD BOOK",
            font=("Arial", 22, "bold")
        ).pack(pady=20)

        frame = tk.Frame(self.root)
        frame.pack()

        labels = [
            "Book ID",
            "Title",
            "Author",
            "Category",
            "Price",
            "Quantity"
        ]

        self.add_entries = {}

        for label in labels:

            tk.Label(
                frame,
                text=label,
                font=("Arial", 11)
            ).pack(anchor="w", pady=(7, 2))

            entry = tk.Entry(
                frame,
                font=("Arial", 12),
                width=30
            )

            entry.pack()

            self.add_entries[label] = entry

        tk.Button(
            self.root,
            text="SAVE BOOK",
            font=("Arial", 12, "bold"),
            width=20,
            height=2,
            command=self.save_new_book
        ).pack(pady=20)

        tk.Button(
            self.root,
            text="BACK",
            width=15,
            command=self.show_dashboard
        ).pack()


    def save_new_book(self):

        book_id = self.add_entries["Book ID"].get().strip()
        title = self.add_entries["Title"].get().strip()
        author = self.add_entries["Author"].get().strip()
        category = self.add_entries["Category"].get().strip()
        price = self.add_entries["Price"].get().strip()
        quantity = self.add_entries["Quantity"].get().strip()

        if not book_id or not title or not author:

            messagebox.showwarning(
                "Missing Data",
                "Book ID, Title and Author are required."
            )

            return

        try:
            price = float(price)
            quantity = int(quantity)

        except:

            messagebox.showerror(
                "Invalid Data",
                "Price must be a number and Quantity must be an integer."
            )

            return

        books = load_books()

        # Check duplicate ID

        for book in books:

            if str(book[0]) == book_id:

                messagebox.showerror(
                    "Duplicate ID",
                    "This Book ID already exists."
                )

                return

        books.append([
            book_id,
            title,
            author,
            category,
            price,
            quantity
        ])

        save_books(books)

        messagebox.showinfo(
            "Success",
            "Book added successfully."
        )

        self.show_dashboard()


    # =====================================================
    # VIEW BOOKS
    # =====================================================

    def view_books(self):

        self.clear_screen()

        tk.Label(
            self.root,
            text="ALL BOOKS",
            font=("Arial", 20, "bold")
        ).pack(pady=10)

        frame = tk.Frame(self.root)
        frame.pack(fill="both", expand=True, padx=5)

        columns = (
            "ID",
            "Title",
            "Author",
            "Category",
            "Price",
            "Qty"
        )

        self.tree = ttk.Treeview(
            frame,
            columns=columns,
            show="headings",
            height=22
        )

        for col in columns:

            self.tree.heading(
                col,
                text=col
            )

        self.tree.column("ID", width=55)
        self.tree.column("Title", width=100)
        self.tree.column("Author", width=100)
        self.tree.column("Category", width=90)
        self.tree.column("Price", width=65)
        self.tree.column("Qty", width=50)

        self.tree.pack(
            side="left",
            fill="both",
            expand=True
        )

        scrollbar = ttk.Scrollbar(
            frame,
            orient="vertical",
            command=self.tree.yview
        )

        scrollbar.pack(
            side="right",
            fill="y"
        )

        self.tree.configure(
            yscrollcommand=scrollbar.set
        )

        books = load_books()

        for book in books:

            self.tree.insert(
                "",
                "end",
                values=book
            )

        tk.Button(
            self.root,
            text="BACK",
            width=15,
            command=self.show_dashboard
        ).pack(pady=10)


    # =====================================================
    # SEARCH BOOK
    # =====================================================

    def search_book(self):

        self.clear_screen()

        tk.Label(
            self.root,
            text="SEARCH BOOK",
            font=("Arial", 20, "bold")
        ).pack(pady=20)

        tk.Label(
            self.root,
            text="Enter Book ID / Title / Author",
            font=("Arial", 11)
        ).pack()

        self.search_entry = tk.Entry(
            self.root,
            font=("Arial", 13),
            width=30
        )

        self.search_entry.pack(pady=10)

        tk.Button(
            self.root,
            text="SEARCH",
            font=("Arial", 12, "bold"),
            width=18,
            height=2,
            command=self.perform_search
        ).pack(pady=10)

        self.search_result = tk.Text(
            self.root,
            width=43,
            height=22,
            font=("Arial", 10)
        )

        self.search_result.pack(pady=10)

        tk.Button(
            self.root,
            text="BACK",
            width=15,
            command=self.show_dashboard
        ).pack()


    def perform_search(self):

        keyword = self.search_entry.get().strip().lower()

        self.search_result.delete(
            "1.0",
            tk.END
        )

        if not keyword:

            self.search_result.insert(
                tk.END,
                "Please enter something to search."
            )

            return

        books = load_books()

        found = False

        for book in books:

            book_id = str(book[0]).lower()
            title = str(book[1]).lower()
            author = str(book[2]).lower()

            if (
                keyword in book_id
                or keyword in title
                or keyword in author
            ):

                found = True

                self.search_result.insert(
                    tk.END,
                    "Book ID: " + str(book[0]) + "\n"
                    "Title: " + str(book[1]) + "\n"
                    "Author: " + str(book[2]) + "\n"
                    "Category: " + str(book[3]) + "\n"
                    "Price: ₹" + str(book[4]) + "\n"
                    "Quantity: " + str(book[5]) + "\n"
                    "-----------------------------\n"
                )

        if not found:

            self.search_result.insert(
                tk.END,
                "No book found."
            )


    # =====================================================
    # UPDATE BOOK
    # =====================================================

    def update_book(self):

        self.clear_screen()

        tk.Label(
            self.root,
            text="UPDATE BOOK",
            font=("Arial", 20, "bold")
        ).pack(pady=15)

        tk.Label(
            self.root,
            text="Enter Book ID",
            font=("Arial", 11)
        ).pack()

        self.update_id = tk.Entry(
            self.root,
            font=("Arial", 13),
            width=25
        )

        self.update_id.pack(pady=8)

        tk.Button(
            self.root,
            text="LOAD BOOK",
            font=("Arial", 11, "bold"),
            command=self.load_book_for_update
        ).pack(pady=8)

        self.update_frame = tk.Frame(
            self.root
        )

        self.update_frame.pack(pady=10)

        tk.Button(
            self.root,
            text="BACK",
            width=15,
            command=self.show_dashboard
        ).pack(pady=10)


    def load_book_for_update(self):

        book_id = self.update_id.get().strip()

        books = load_books()

        selected = None

        for book in books:

            if str(book[0]) == book_id:

                selected = book
                break

        if selected is None:

            messagebox.showerror(
                "Not Found",
                "Book ID not found."
            )

            return

        for widget in self.update_frame.winfo_children():
            widget.destroy()

        labels = [
            "Title",
            "Author",
            "Category",
            "Price",
            "Quantity"
        ]

        self.update_entries = {}

        values = selected[1:]

        for i in range(len(labels)):

            tk.Label(
                self.update_frame,
                text=labels[i]
            ).grid(
                row=i,
                column=0,
                padx=5,
                pady=5
            )

            entry = tk.Entry(
                self.update_frame,
                width=22
            )

            entry.insert(
                0,
                str(values[i])
            )

            entry.grid(
                row=i,
                column=1,
                padx=5,
                pady=5
            )

            self.update_entries[labels[i]] = entry

        tk.Button(
            self.update_frame,
            text="UPDATE",
            font=("Arial", 11, "bold"),
            command=lambda: self.save_update(book_id)
        ).grid(
            row=5,
            column=0,
            columnspan=2,
            pady=15
        )


    def save_update(self, book_id):

        try:

            title = self.update_entries["Title"].get()
            author = self.update_entries["Author"].get()
            category = self.update_entries["Category"].get()
            price = float(
                self.update_entries["Price"].get()
            )
            quantity = int(
                self.update_entries["Quantity"].get()
            )

        except:

            messagebox.showerror(
                "Invalid Data",
                "Please enter valid price and quantity."
            )

            return

        books = load_books()

        for i in range(len(books)):

            if str(books[i][0]) == book_id:

                books[i] = [
                    book_id,
                    title,
                    author,
                    category,
                    price,
                    quantity
                ]

                break

        save_books(books)

        messagebox.showinfo(
            "Success",
            "Book updated successfully."
        )

        self.show_dashboard()


    # =====================================================
    # DELETE BOOK
    # =====================================================

    def delete_book(self):

        self.clear_screen()

        tk.Label(
            self.root,
            text="DELETE BOOK",
            font=("Arial", 20, "bold")
        ).pack(pady=30)

        tk.Label(
            self.root,
            text="Enter Book ID",
            font=("Arial", 12)
        ).pack()

        self.delete_id = tk.Entry(
            self.root,
            font=("Arial", 13),
            width=25
        )

        self.delete_id.pack(pady=10)

        tk.Button(
            self.root,
            text="DELETE",
            font=("Arial", 12, "bold"),
            width=18,
            height=2,
            command=self.perform_delete
        ).pack(pady=15)

        tk.Button(
            self.root,
            text="BACK",
            width=15,
            command=self.show_dashboard
        ).pack()


    def perform_delete(self):

        book_id = self.delete_id.get().strip()

        if not book_id:

            messagebox.showwarning(
                "Input Required",
                "Enter Book ID."
            )

            return

        books = load_books()

        found = False

        new_books = []

        for book in books:

            if str(book[0]) == book_id:

                found = True

            else:

                new_books.append(book)

        if not found:

            messagebox.showerror(
                "Not Found",
                "Book ID not found."
            )

            return

        confirm = messagebox.askyesno(
            "Confirm Delete",
            "Are you sure you want to delete this book?"
        )

        if confirm:

            save_books(new_books)

            messagebox.showinfo(
                "Success",
                "Book deleted successfully."
            )

            self.show_dashboard()


# =========================================================
# START PROGRAM
# =========================================================

if __name__ == "__main__":

    create_database()

    root = tk.Tk()

    app = BookInventoryApp(root)

    root.mainloop()